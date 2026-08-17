import json
import logging
import sys
import tomllib
from argparse import Namespace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest

from cofr.domain import Claim, OpenQuestion, Risk, validate_and_normalize
from cofr.ingest import extract_pdf_text
from cofr.cli import cmd_refresh
from cofr.renames import _pre_load_pending_pack_fixup, load_pending_renames
from cofr.state import CorruptStateError, apply_parsed_records, init, load_config, load_index, load_state, migrate_rollback
from cofr.synthesis import compute_computed_risks
from conftest import run_cofr


def test_import_metadata_name_matches_distribution_name():
    import cofr

    pyproject = tomllib.loads((Path(__file__).parents[1] / 'pyproject.toml').read_text())
    assert cofr.DISTRIBUTION_NAME == pyproject['project']['name']


def test_readme_identifies_cofr_as_gist_implementation():
    readme = (Path(__file__).parents[1] / 'README.md').read_text()
    assert 'concrete implementation of the [LLM Research Manager pattern]' in readme
    assert '990a26d214c94261fcf10c0506bfa156' in readme


def test_readme_is_prompt_first_and_moves_cli_details_to_technical_docs():
    readme = (Path(__file__).parents[1] / 'README.md').read_text()
    quick_start = readme.split('## Quick start', 1)[1].split('## Technical documentation', 1)[0]
    assert 'Paste this prompt' in quick_start
    assert '```bash' not in quick_start
    assert 'cofr --help' not in quick_start
    assert 'workflow' not in quick_start.lower()
    assert ' in COFR' not in quick_start
    assert ' into COFR' not in quick_start
    assert "Based on COFR's outline of the project, create an HTML overview of the current project state." in quick_start
    assert '[Technical reference](https://github.com/GPTomics/cofr/blob/main/tech_docs.md)' in readme
    assert '### Commands' not in readme
    assert 'pip install git+https://github.com/GPTomics/cofr.git' in readme
    assert 'pipx install cofr-research' not in readme
    assert 'uv tool install cofr-research' not in readme
    assert 'COFR itself does not ship or require agent skills' in readme
    prompts = [part.split('```', 1)[0].strip() for part in quick_start.split('```text')[1:]]
    assert prompts
    assert max(map(len, prompts)) < 200

    technical_docs = (Path(__file__).parents[1] / 'tech_docs.md').read_text()
    assert '## Commands' in technical_docs
    assert '## Object types' in technical_docs
    assert '## JSON envelope' in technical_docs
    assert 'uvx --from cofr-research' not in technical_docs

    help_text = (Path(__file__).parents[1] / 'src' / 'cofr' / 'help.txt').read_text()
    assert 'Authoring depth:' in help_text
    assert 'Preserve technical depth whenever the source material provides it.' in help_text


def test_sdist_excludes_repository_only_material():
    pyproject = tomllib.loads((Path(__file__).parents[1] / 'pyproject.toml').read_text())
    excluded = set(pyproject['tool']['hatch']['build']['targets']['sdist']['exclude'])
    assert {
        '/.claude', '/AGENTS.md', '/CLAUDE.md', '/ARCHITECTURE_PLAN.md',
        '/M1_SPEC.md', '/bugs.md',
    } <= excluded


def test_obsolete_product_outline_is_not_tracked_as_project_documentation():
    assert not (Path(__file__).parents[1] / 'instructions.md').exists()


def test_deleted_claim_is_retired_and_timeline_records_transition():
    state = {
        'claims': [], 'evidence': [], 'experiments': [], 'decisions': [],
        'open_questions': [], 'risks': [], 'artifacts': [],
    }
    claim = Claim(id='claim_a', status='supported', parsed_from='claims.yaml#claim_a')
    apply_parsed_records(
        state, [claim], {'claims.yaml': {}},
        packs_parsed_successfully={'claims.yaml'},
    )

    diff = apply_parsed_records(
        state, [], {'claims.yaml': {}},
        packs_parsed_successfully={'claims.yaml'},
    )

    retired = state['claims'][0]
    assert retired['status'] == 'retired'
    assert retired['stale'] is False
    assert retired['source_missing'] is True
    assert retired['_status_timeline'][-1]['s'] == 'retired'
    assert diff['stale'] == [{'type': 'claim', 'id': 'claim_a'}]


def test_user_authored_computed_risk_source_is_normalized(tmp_path):
    run_cofr(['init', str(tmp_path)])
    (tmp_path / 'risks.yaml').write_text(
        '- type: risk\n  id: risk_fake\n  source: computed\n  statement: User input\n',
        encoding='utf-8',
    )

    result = run_cofr(['refresh', '--json', str(tmp_path)])

    assert result.returncode == 1
    payload = json.loads(result.stdout)
    assert payload['data']['state']['risks'][0]['source'] == 'authored'
    warning_text = ' '.join(payload['data']['warnings']).lower()
    assert 'reserved' in warning_text and 'authored' in warning_text


def test_unknown_binary_yaml_value_cannot_break_refresh(tmp_path):
    run_cofr(['init', str(tmp_path)])
    record = '''---
type: claim
id: claim_binary
custom_blob: !!binary aGVsbG8=
custom_score: .nan
---

## Statement

Safe state must still be serializable.
'''

    result = run_cofr(['add', str(tmp_path)], stdin=record)

    assert result.returncode in (0, 1)
    state = json.loads((tmp_path / '.cofr' / 'state.json').read_text(encoding='utf-8'))
    assert state['claims'][0]['id'] == 'claim_binary'
    assert state['claims'][0]['_unknown_fields']['custom_score'] is None
    json.dumps(state)


def test_malformed_nested_fields_warn_and_do_not_crash_refresh(tmp_path):
    run_cofr(['init', str(tmp_path)])
    (tmp_path / 'evidences').mkdir()
    (tmp_path / 'evidences' / 'manual.yaml').write_text(
        '- type: evidence\n  id: ev_bad\n  claim_links: wrong-shape\n'
        '  source_anchors: wrong-shape\n',
        encoding='utf-8',
    )

    result = run_cofr(['refresh', '--json', str(tmp_path)])

    assert result.returncode == 1
    payload = json.loads(result.stdout)
    evidence = payload['data']['state']['evidence'][0]
    assert evidence['claim_links'] == []
    assert evidence['source_anchors'] == []


def test_corrupt_index_is_reported_as_corrupt_state(tmp_path):
    init(tmp_path)
    (tmp_path / '.cofr' / 'index.json').write_text('{bad json', encoding='utf-8')

    with pytest.raises(CorruptStateError, match='index.json'):
        load_index(tmp_path)
    result = run_cofr(['show', 'state', str(tmp_path)])
    assert result.returncode == 4
    assert 'traceback' not in result.stderr.lower()


def test_non_utf8_state_is_reported_without_traceback(tmp_path):
    init(tmp_path)
    (tmp_path / '.cofr' / 'state.json').write_bytes(b'\xff\xfe')

    result = run_cofr(['show', 'state', str(tmp_path)])

    assert result.returncode == 4
    assert 'traceback' not in result.stderr.lower()


def test_pending_rename_rejects_pack_path_outside_project(tmp_path):
    init(tmp_path)
    outside = tmp_path.parent / 'outside.yaml'
    outside.write_text('- type: claim\n  id: old\n', encoding='utf-8')
    pending_path = tmp_path / '.cofr' / 'pending_renames.json'
    pending_path.write_text(json.dumps({
        'schema_version': 1,
        'entries': [{
            'type': 'claim', 'old_id': 'old', 'new_id': 'new',
            'pack_path': '../outside.yaml', 'mode': 'standard',
        }],
    }), encoding='utf-8')

    assert load_pending_renames(tmp_path) is None
    result = _pre_load_pending_pack_fixup(tmp_path, json.loads(pending_path.read_text()))
    assert result['action'] == 'refuse'
    assert outside.read_text(encoding='utf-8').startswith('- type: claim')


def test_migration_rollback_rejects_manifest_paths_outside_project(tmp_path):
    init(tmp_path)
    outside = tmp_path.parent / 'outside-pack.yaml'
    outside.write_text('do not move', encoding='utf-8')
    snapshot = tmp_path / '.cofr' / 'history' / 'safe.json'
    snapshot.write_text((tmp_path / '.cofr' / 'state.json').read_text(), encoding='utf-8')
    (tmp_path / '.cofr' / 'migration_manifest.json').write_text(json.dumps({
        'operations': [{
            'kind': 'commit_pack', 'src': '.cofr/migrate_pending/x.yaml',
            'dst': '../outside-pack.yaml', 'status': 'done',
        }],
    }), encoding='utf-8')

    with pytest.raises(RuntimeError, match='outside project'):
        migrate_rollback(tmp_path, from_history=snapshot, confirm=True)
    assert outside.read_text(encoding='utf-8') == 'do not move'


def test_computed_risk_collision_with_authored_id_is_dropped():
    state = {
        'claims': [], 'evidence': [], 'experiments': [], 'decisions': [],
        'open_questions': [], 'artifacts': [],
        'risks': [Risk(
            id='risk_computed_orphaned_assumption_claim_a',
            statement='Authored record',
        ).__dict__],
    }
    contradictions = {
        'claim_unchanged': [], 'decision_basis_eroded': [],
        'evidence_conflict': [], 'eroded_confidence': [],
        'orphaned_assumption': [{
            'claim_id': 'claim_a', 'reason': 'collision', 'cited_ids': ['claim_a'],
        }],
    }

    computed, warnings = compute_computed_risks(state, contradictions=contradictions)

    assert computed == []
    assert any('authored' in warning and 'collision' in warning for warning in warnings)


def test_m4_image_docx_and_xlsx_indexing_and_mentions(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image

    run_cofr(['init', str(tmp_path)])
    (tmp_path / 'claims.yaml').write_text(
        '- type: claim\n  id: claim_alpha\n  title: Alpha\n', encoding='utf-8',
    )
    Image.new('RGB', (13, 17), color='navy').save(tmp_path / 'figure.png')
    document = Document()
    document.add_paragraph('Document discussion of claim_alpha.')
    document.save(tmp_path / 'notes.docx')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Results'
    sheet.append(['claim', 'note'])
    sheet.append(['claim_alpha', 'measured'])
    workbook.save(tmp_path / 'results.xlsx')

    result = run_cofr(['refresh', '--json', str(tmp_path)])

    assert result.returncode in (0, 1), result.stderr
    index = json.loads((tmp_path / '.cofr' / 'index.json').read_text())
    assert index['figure.png']['classification'] == 'content_extracted'
    assert index['figure.png']['image_dimensions'] == [13, 17]
    assert index['figure.png']['image_format'] == 'PNG'
    assert index['figure.png']['image_exif'] == {}
    assert index['figure.png']['extracted_text_length'] == 0
    assert index['notes.docx']['classification'] == 'content_extracted'
    assert index['notes.docx']['extracted_text_length'] > 0
    assert index['results.xlsx']['classification'] == 'content_extracted'
    assert index['results.xlsx']['sheet_names'] == ['Results']
    claims = json.loads(run_cofr(['show', 'claims', '--json', str(tmp_path)]).stdout)
    assert claims['data']['claims'][0]['mentioned_in'] == ['notes.docx', 'results.xlsx']


def test_opaque_non_text_files_do_not_get_content_hash(tmp_path):
    run_cofr(['init', str(tmp_path)])
    (tmp_path / 'opaque.bin').write_bytes(b'\x00\xffopaque')

    result = run_cofr(['refresh', '--json', str(tmp_path)])

    assert result.returncode in (0, 1)
    index = json.loads((tmp_path / '.cofr' / 'index.json').read_text())
    assert index['opaque.bin']['classification'] == 'non_text'
    assert index['opaque.bin']['content_hash'] is None
    assert index['opaque.bin']['extracted_text_length'] is None


def test_init_defaults_project_name_and_reports_idempotence(tmp_path):
    first = run_cofr(['init', str(tmp_path)])
    second = run_cofr(['init', str(tmp_path)])

    assert first.returncode == second.returncode == 0
    assert f'project_name: {tmp_path.name}' in (tmp_path / '.cofr' / 'config.yaml').read_text()
    assert 'already initialized' in second.stdout.lower()


def test_show_state_human_output_is_full_state_not_only_counts(tmp_path):
    run_cofr(['init', str(tmp_path)])
    (tmp_path / 'claims.yaml').write_text(
        '- type: claim\n  id: claim_full\n  title: Full record\n  statement: Visible prose\n',
        encoding='utf-8',
    )
    run_cofr(['refresh', str(tmp_path)])

    result = run_cofr(['show', 'state', str(tmp_path)])

    assert result.returncode == 0
    assert 'claim_full' in result.stdout
    assert 'Visible prose' in result.stdout


def test_persisted_ids_must_be_valid_and_globally_unique(tmp_path):
    init(tmp_path)
    state_path = tmp_path / '.cofr' / 'state.json'
    state = json.loads(state_path.read_text())
    state['claims'] = [{'id': 'same'}]
    state['risks'] = [{'id': 'same'}]
    state_path.write_text(json.dumps(state), encoding='utf-8')

    with pytest.raises(CorruptStateError, match='duplicate id'):
        load_state(tmp_path)


def test_pdfminer_internal_warnings_do_not_pollute_cli_stderr(tmp_path, monkeypatch, caplog):
    pdf = tmp_path / 'minimal.pdf'
    pdf.write_bytes(b'%PDF-1.4\n')

    class FakeMarkItDown:
        def convert(self, _path):
            logging.getLogger('pdfminer.pdfdocument').warning('internal parser noise')
            return SimpleNamespace(text_content='extracted')

    monkeypatch.setitem(sys.modules, 'markitdown', SimpleNamespace(MarkItDown=FakeMarkItDown))
    with caplog.at_level(logging.WARNING):
        text, _pages = extract_pdf_text(pdf)

    assert text == 'extracted'
    assert 'internal parser noise' not in caplog.text


def test_non_scalar_question_priority_normalizes_without_crashing():
    question = OpenQuestion(id='q_shape', priority={'invalid': 'mapping'})

    normalized, warnings = validate_and_normalize(question)

    assert normalized.priority == 'medium'
    assert normalized.blocking_severity == 'medium'
    assert warnings


def test_config_text_fields_reject_nested_values(tmp_path):
    init(tmp_path)
    (tmp_path / '.cofr' / 'config.yaml').write_text(
        'project_name: [not, text]\nproject_objective: {not: text}\n',
        encoding='utf-8',
    )

    config, warnings = load_config(tmp_path)

    assert config['project_name'] == ''
    assert config['project_objective'] == ''
    assert len(warnings) == 2


def test_refresh_write_failure_returns_controlled_corrupt_exit(tmp_path, capsys):
    init(tmp_path)
    with patch('cofr.cli.save_state', side_effect=OSError('disk unavailable')):
        result = cmd_refresh(Namespace(
            project_path=str(tmp_path), json=False,
            rebuild_timelines=False, rebuild_renames_log=False,
        ))

    assert result == 4
    assert 'disk unavailable' in capsys.readouterr().err
